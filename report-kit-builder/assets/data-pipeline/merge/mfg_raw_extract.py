# -*- coding: utf-8 -*-
"""制造车间 raw_reading 全序列提取 → 分片 INSERT SQL（mfg_raw_reading: tag/ts/val）
源：
  A. 云平台 zip_xlsx（从机格式: col0=时间 col2=冷量 col3=电量，30s 级）
  B. WinCC 导出「累计冷量*.csv」「累计用电量*.csv」「能效报表*.csv」（能效报表含累计制冷量+累计用电量两列）
排除路径含「配料」的子目录（制造 rar 里混入了配料数据夹）。
"""
import csv, glob, os, zipfile, warnings
warnings.filterwarnings('ignore')
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # 工作空间根，换机/改目录自动适应
BASE = os.path.join(ROOT, '03_processed', '制造', '.cache_rar')
OUT = os.path.join(ROOT, '08_scratch', 'mfg_raw_chunks')

def is_mfg(fp):
    return '配料' not in fp

def num(v):
    try:
        f = float(v)
        return f if f == f else None
    except (TypeError, ValueError):
        return None

def norm_ts(t):
    t = str(t)[:19].replace('/', '-').replace('T', ' ')
    return t if len(t) >= 16 else None

rows = {}  # (tag, ts) -> val

# A. 云平台 zip_xlsx
for fp in glob.glob(os.path.join(BASE, '制造车间云平台历史数据', 'zip_*', '*.xlsx')):
    try:
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    except Exception as e:
        print('skip', fp, e); continue
    ws = wb[wb.sheetnames[0]]
    n = 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None:
            continue
        ts = norm_ts(r[0])
        if not ts:
            continue
        c, e = num(r[2] if len(r) > 2 else None), num(r[3] if len(r) > 3 else None)
        if c is not None: rows[('cool', ts)] = c; n += 1
        if e is not None: rows[('elec', ts)] = e; n += 1
    wb.close()
    print('usr:', os.path.basename(fp), n)

# B. WinCC CSV（累计冷量/累计用电量/能效报表）
def read_any(fp):
    for enc in ('utf-16', 'utf-8-sig', 'gbk'):
        try:
            return open(fp, encoding=enc, newline='').read()
        except Exception:
            continue
    return None

for fp in glob.glob(os.path.join(BASE, '制造车间数据1-9月', '**', '*.csv'), recursive=True):
    if not is_mfg(fp):
        continue
    name = os.path.basename(fp)
    txt = read_any(fp)
    if txt is None:
        continue
    sep = '\t' if '\t' in txt.splitlines()[1 if len(txt.splitlines()) > 1 else 0] else ','
    n = 0
    lines = txt.splitlines()[1:]
    for l in lines:
        c = l.split(sep)
        if len(c) < 2 or not c[0].strip():
            continue
        ts = norm_ts(c[0].strip())
        if not ts:
            continue
        if '累计冷量' in name:
            v = num(c[1])
            if v is not None: rows[('cool', ts)] = v; n += 1
        elif '累计用电量' in name:
            v = num(c[1])
            if v is not None: rows[('elec', ts)] = v; n += 1
        elif '能效报表' in name and len(c) >= 4:
            cool, elec = num(c[2]), num(c[3])
            if cool is not None: rows[('cool', ts)] = cool; n += 1
            if elec is not None: rows[('elec', ts)] = elec; n += 1
    if n:
        print('wincc:', name, n)

items = sorted(rows.items(), key=lambda kv: (kv[0][1], kv[0][0]))
print('total rows:', len(items))
os.makedirs(OUT, exist_ok=True)
CH = 1000
files = []
for i in range(0, len(items), CH):
    chunk = items[i:i + CH]
    vals = ', '.join("('%s', '%s', %s)" % (t, ts, v) for (t, ts), v in chunk)
    fn = os.path.join(OUT, 'c%04d.sql' % (i // CH))
    open(fn, 'w', encoding='utf-8').write(
        "INSERT INTO workspace_aadkvj7vniyyw.mfg_raw_reading (tag, ts, val) VALUES " + vals + ";\n")
    files.append(fn)
print('chunks:', len(files))
